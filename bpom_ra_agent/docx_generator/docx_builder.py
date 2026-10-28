"""Word Document (.docx) Report Builder for BPOM Regulatory Submissions.
Produces executive-grade, beautifully formatted Word dossiers, checklists,
and gap audit reports with styled tables, callout blocks, and bilingual labels.
"""

import os
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Color Palette
COLOR_PRIMARY_NAVY = RGBColor(26, 54, 93)     # #1A365D
COLOR_SECONDARY_BLUE = RGBColor(43, 108, 176) # #2B6CB0
COLOR_DARK_TEXT = RGBColor(45, 55, 72)        # #2D3748
COLOR_MUTED = RGBColor(113, 128, 150)         # #718096
COLOR_RED_CRITICAL = RGBColor(197, 48, 48)    # #C53030
COLOR_AMBER_MAJOR = RGBColor(183, 121, 31)    # #B7791F
COLOR_GREEN_OK = RGBColor(40, 116, 75)        # #28744B

HEX_HEADER_NAVY = "1A365D"
HEX_LIGHT_GRAY = "F7FAFC"
HEX_BORDER_GRAY = "CBD5E0"
HEX_CRITICAL_BG = "FFF5F5"
HEX_MAJOR_BG = "FFFAF0"
HEX_OK_BG = "F0FFF4"

def set_cell_background(cell, hex_color: str):
    """Sets background shading of a docx table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a docx cell in twips."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def style_row_as_header(row, col_widths: List[float]):
    for i, cell in enumerate(row.cells):
        set_cell_background(cell, HEX_HEADER_NAVY)
        set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
        cell.width = Inches(col_widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(10)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)


class DocxReportBuilder:
    """Builds professional regulatory assessment and audit documents."""

    @staticmethod
    def create_evaluation_document(assessment: Any, product: Dict[str, Any],
                                   checklist: Dict[str, Any], language: str,
                                   output_path: str):
        """Generates comprehensive Drug Registration Pathway & ACTD Dossier Document."""
        doc = Document()
        is_id = (language.upper() == "ID")

        # Page setup
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
            
            # Header
            header_p = section.header.paragraphs[0]
            header_p.text = "BADAN POM RI - REGULATORY AFFAIRS DOSSIER ASSESSMENT REPORT"
            header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            header_p.runs[0].font.size = Pt(8.5)
            header_p.runs[0].font.color.rgb = COLOR_MUTED

            # Footer
            footer_p = section.footer.paragraphs[0]
            footer_p.text = "CONFIDENTIAL - Strictly for Indonesian Drug Registration & Dossier Submission Preparation"
            footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_p.runs[0].font.size = Pt(8)
            footer_p.runs[0].font.color.rgb = COLOR_MUTED

        # 1. Main Header Title
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(4)
        run_title = title_p.add_run(
            "LAPORAN PENILAIAN JALUR REGISTRASI OBAT & KESIAPAN ACTD" if is_id
            else "DRUG REGISTRATION PATHWAY & ACTD COMPLIANCE ASSESSMENT REPORT"
        )
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_PRIMARY_NAVY

        # Subtitle
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(16)
        run_sub = sub_p.add_run(
            f"Evaluasi Regulasi Berdasarkan Peraturan BPOM No. 24/2017 jo. PerBPOM 13/2021 & 24/2023 | Tarif PNBP PP 32/2017" if is_id
            else f"Regulatory Evaluation under BPOM Reg No. 24/2017 jo. Reg 13/2021 & 24/2023 | Non-Tax State Revenue PP 32/2017"
        )
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(10.5)
        run_sub.font.color.rgb = COLOR_SECONDARY_BLUE

        # 2. Executive Product Profile Table
        doc.add_heading("1. Profil Produk & Parameter Pengajuan" if is_id else "1. Product Profile & Submission Parameters", level=2)
        
        table = doc.add_table(rows=6, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [2.2, 4.8]
        
        rows_data = [
            ("Nama Produk (Brand Name)" if is_id else "Product Brand Name", product.get("product_name", "-")),
            ("Zat Aktif (INN)" if is_id else "Active Substance (INN)", product.get("active_substance", "-")),
            ("Bentuk Sediaan & Kekuatan" if is_id else "Dosage Form & Strength", f"{product.get('dosage_form', '-')} | {product.get('strength', '-')}"),
            ("Klasifikasi Molekul" if is_id else "Molecule Classification", product.get("product_type", "-")),
            ("Asal Pembuatan (Origin)" if is_id else "Origin & Site Status", f"{product.get('origin', '-')} ({product.get('foreign_gmp_status', 'LOCAL_CPOB')})"),
            ("Status Persetujuan Internasional" if is_id else "International Approval Status", 
             f"SRA Approved: {product.get('is_sra_approved', False)} ({', '.join(product.get('sra_countries', [])) if product.get('sra_countries') else 'None'})")
        ]

        for idx, (label, val) in enumerate(rows_data):
            row = table.rows[idx]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width, c1.width = Inches(widths[0]), Inches(widths[1])
            set_cell_background(c0, HEX_LIGHT_GRAY)
            set_cell_margins(c0, 80, 80, 100, 100)
            set_cell_margins(c1, 80, 80, 100, 100)
            
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(label)
            r0.font.bold = True
            r0.font.size = Pt(9.5)
            r0.font.color.rgb = COLOR_DARK_TEXT

            p1 = c1.paragraphs[0]
            r1 = p1.add_run(str(val))
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = COLOR_DARK_TEXT

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 3. Regulatory Decision Matrix (Category, Track, Timelines, Fees)
        doc.add_heading("2. Hasil Evaluasi Jalur, Biaya PNBP & Timeline Evaluasi BPOM" if is_id else "2. BPOM Pathway Determination, PNBP Fees & Timelines", level=2)

        callout = doc.add_table(rows=1, cols=1)
        callout.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_cell = callout.rows[0].cells[0]
        c_cell.width = Inches(7.0)
        set_cell_background(c_cell, "EBF8FF")
        set_cell_margins(c_cell, 120, 120, 150, 150)
        cp = c_cell.paragraphs[0]
        cr = cp.add_run(
            f"KATEGORI REGISTRASI: {assessment.category_name_id}\n"
            f"JALUR EVALUASI: {assessment.track_name_id}\n"
            f"TOTAL BIAYA RESMI PNBP (PP 32/2017): Rp {assessment.total_fee_idr:,.0f}\n"
            f"ESTIMASI WAKTU EVALUASI: {assessment.prereg_days_hk} HK Pra-Registrasi + {assessment.timeline_days_hk} HK Registrasi = {assessment.total_working_days} Hari Kerja (~{assessment.estimated_calendar_months} Bulan Kalender)"
            if is_id else
            f"REGISTRATION CATEGORY: {assessment.category_name_en}\n"
            f"EVALUATION TRACK: {assessment.track_name_en}\n"
            f"TOTAL STATUTORY PNBP FEES (PP 32/2017): IDR {assessment.total_fee_idr:,.0f}\n"
            f"REVIEW TIMELINE: {assessment.prereg_days_hk} WD Pre-Reg + {assessment.timeline_days_hk} WD Dossier = {assessment.total_working_days} Working Days (~{assessment.estimated_calendar_months} Calendar Months)"
        )
        cr.font.name = "Calibri"
        cr.font.size = Pt(10)
        cr.font.bold = True
        cr.font.color.rgb = COLOR_PRIMARY_NAVY

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # PNBP Fees Breakdown Table
        doc.add_paragraph("Rincian Biaya Tarif PNBP Sesuai PP No. 32 Tahun 2017:" if is_id else "Breakdown of Official Non-Tax State Revenue (PNBP) Tariffs under PP 32/2017:")
        fee_table = doc.add_table(rows=len(assessment.fees) + 2, cols=4)
        fee_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        fee_widths = [1.2, 3.2, 1.4, 1.2]
        
        # Header
        h_row = fee_table.rows[0]
        h_row.cells[0].paragraphs[0].text = "Kode Tarif" if is_id else "Tariff Code"
        h_row.cells[1].paragraphs[0].text = "Uraian Biaya PNBP" if is_id else "Fee Description"
        h_row.cells[2].paragraphs[0].text = "Dasar Hukum" if is_id else "Legal Basis"
        h_row.cells[3].paragraphs[0].text = "Nominal (IDR)" if is_id else "Amount (IDR)"
        style_row_as_header(h_row, fee_widths)

        for i, fee in enumerate(assessment.fees):
            row = fee_table.rows[i + 1]
            row.cells[0].paragraphs[0].text = fee["code"]
            row.cells[1].paragraphs[0].text = fee["component"]
            row.cells[2].paragraphs[0].text = fee["basis"]
            row.cells[3].paragraphs[0].text = f"Rp {fee['amount_idr']:,.0f}"
            for c_idx, cell in enumerate(row.cells):
                cell.width = Inches(fee_widths[c_idx])
                set_cell_margins(cell, 80, 80, 100, 100)
                if i % 2 == 1:
                    set_cell_background(cell, HEX_LIGHT_GRAY)
                cell.paragraphs[0].runs[0].font.size = Pt(9)
                cell.paragraphs[0].runs[0].font.color.rgb = COLOR_DARK_TEXT
                if c_idx == 3:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # Total Row
        tot_row = fee_table.rows[-1]
        tot_row.cells[0].paragraphs[0].text = "TOTAL"
        tot_row.cells[1].paragraphs[0].text = "Total Biaya PNBP Wajib Bayar (Billing SIMPONI)" if is_id else "Total Mandatory PNBP Payable (SIMPONI)"
        tot_row.cells[2].paragraphs[0].text = "PP 32/2017"
        tot_row.cells[3].paragraphs[0].text = f"Rp {assessment.total_fee_idr:,.0f}"
        for c_idx, cell in enumerate(tot_row.cells):
            cell.width = Inches(fee_widths[c_idx])
            set_cell_background(cell, "E2E8F0")
            set_cell_margins(cell, 90, 90, 100, 100)
            cell.paragraphs[0].runs[0].font.size = Pt(9.5)
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = COLOR_PRIMARY_NAVY
            if c_idx == 3:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 4. Key Submission Requirements & Risk Watchlist
        doc.add_heading("3. Persyaratan Kritis & Identifikasi Risiko Regulasi" if is_id else "3. Critical Pre-Requisites & Regulatory Risk Watchlist", level=2)
        
        req_list = assessment.key_requirements_id if is_id else assessment.key_requirements_en
        for req in req_list:
            p = doc.add_paragraph(style='List Bullet')
            r = p.add_run(req)
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK_TEXT

        if assessment.risk_factors:
            doc.add_paragraph()
            for risk in assessment.risk_factors:
                rp = doc.add_paragraph()
                rp.paragraph_format.left_indent = Inches(0.2)
                rr = rp.add_run(f"⚠️ {risk}")
                rr.font.size = Pt(9.5)
                rr.font.bold = True
                rr.font.color.rgb = COLOR_RED_CRITICAL

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 5. ACTD Module-by-Module Checklist
        doc.add_heading("4. Daftar Periksa Kelengkapan Dokumen ACTD (Part I - IV)" if is_id else "4. ACTD Module Dossier Completeness Checklist (Parts I - IV)", level=2)
        
        actd_parts = [
            ("Bagian I: Dokumen Administratif & Penandaan" if is_id else "Part I: Administrative Documents & Labeling", checklist["Part_I_Administrative"]),
            ("Bagian II: Dokumen Mutu (Zat Aktif & Produk Jadi)" if is_id else "Part II: Quality Documents (Drug Substance & Product)", checklist["Part_II_Quality"]),
            ("Bagian III: Dokumen Non-Klinis" if is_id else "Part III: Non-Clinical Documents", checklist["Part_III_NonClinical"]),
            ("Bagian IV: Dokumen Klinis" if is_id else "Part IV: Clinical Documents", checklist["Part_IV_Clinical"]),
        ]

        chk_widths = [0.8, 2.5, 0.9, 2.8]

        for part_title, items in actd_parts:
            doc.add_heading(part_title, level=3)
            tbl = doc.add_table(rows=len(items) + 1, cols=4)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            # Header
            hr = tbl.rows[0]
            hr.cells[0].paragraphs[0].text = "Modul" if is_id else "Module"
            hr.cells[1].paragraphs[0].text = "Judul Dokumen ACTD" if is_id else "ACTD Document Element"
            hr.cells[2].paragraphs[0].text = "Status" if is_id else "Status"
            hr.cells[3].paragraphs[0].text = "Catatan Spesifik BPOM RI" if is_id else "BPOM Specific Requirements"
            style_row_as_header(hr, chk_widths)

            for idx, item in enumerate(items):
                r = tbl.rows[idx + 1]
                r.cells[0].paragraphs[0].text = item["code"]
                r.cells[1].paragraphs[0].text = item["title_id"] if is_id else item["title_en"]
                status_text = "[WAIVED]" if not item["required"] else "[WAJIB]" if is_id else "[MANDATORY]"
                r.cells[2].paragraphs[0].text = status_text
                r.cells[3].paragraphs[0].text = item["notes_id"] if is_id else item["notes_en"]

                for c_idx, cell in enumerate(r.cells):
                    cell.width = Inches(chk_widths[c_idx])
                    set_cell_margins(cell, 60, 60, 80, 80)
                    if idx % 2 == 1:
                        set_cell_background(cell, HEX_LIGHT_GRAY)
                    cell.paragraphs[0].runs[0].font.size = Pt(8.5)
                    if c_idx == 2:
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if not item["required"]:
                            cell.paragraphs[0].runs[0].font.color.rgb = COLOR_MUTED
                        else:
                            cell.paragraphs[0].runs[0].font.bold = True
                            cell.paragraphs[0].runs[0].font.color.rgb = COLOR_PRIMARY_NAVY

            doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # Save document
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path

    @staticmethod
    def create_gap_audit_document(audit_result: Dict[str, Any], product: Dict[str, Any],
                                  language: str, output_path: str):
        """Generates comprehensive Dossier Gap-Check & Risk Audit Report."""
        doc = Document()
        is_id = (language.upper() == "ID")

        # Page setup
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
            
            header_p = section.header.paragraphs[0]
            header_p.text = "BPOM DOSSIER AUDIT & GAP-ANALYSIS REPORT"
            header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            header_p.runs[0].font.size = Pt(8.5)
            header_p.runs[0].font.color.rgb = COLOR_MUTED

            footer_p = section.footer.paragraphs[0]
            footer_p.text = "CONFIDENTIAL REGULATORY AUDIT - Pre-Submission Gap Assessment"
            footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_p.runs[0].font.size = Pt(8)
            footer_p.runs[0].font.color.rgb = COLOR_MUTED

        # Main Title
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(4)
        run_title = title_p.add_run(
            "LAPORAN AUDIT & ANALISIS KESENJANGAN DOSSIER (GAP-CHECK)" if is_id
            else "DOSSIER GAP-CHECK & REGULATORY RISK AUDIT REPORT"
        )
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(20)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_RED_CRITICAL if audit_result["critical_count"] > 0 else COLOR_PRIMARY_NAVY

        # Subtitle
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(16)
        run_sub = sub_p.add_run(
            f"Audit Kepatuhan Format ACTD & Standar Teknis BPOM RI | Produk: {product.get('product_name', '-')}" if is_id
            else f"Compliance Audit against BPOM ACTD Regulations & Guidelines | Target Product: {product.get('product_name', '-')}"
        )
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(10.5)
        run_sub.font.color.rgb = COLOR_SECONDARY_BLUE

        # 1. Executive Summary Callout Box
        status_box = doc.add_table(rows=1, cols=1)
        status_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        sb_cell = status_box.rows[0].cells[0]
        sb_cell.width = Inches(7.0)
        
        bg_color = HEX_CRITICAL_BG if audit_result["critical_count"] > 0 else HEX_MAJOR_BG
        border_text_color = COLOR_RED_CRITICAL if audit_result["critical_count"] > 0 else COLOR_AMBER_MAJOR
        set_cell_background(sb_cell, bg_color)
        set_cell_margins(sb_cell, 120, 120, 150, 150)
        
        sb_p = sb_cell.paragraphs[0]
        sb_run = sb_p.add_run(
            f"STATUS KESIAPAN DOSSIER: {audit_result['status_text_id']}\n"
            f"RINGKASAN TEMUAN: Total {audit_result['total_gaps']} Temuan | "
            f"Kritis (Critical): {audit_result['critical_count']} | "
            f"Mayor (Major): {audit_result['major_count']} | "
            f"Minor: {audit_result['minor_count']}\n"
            f"REKOMENDASI: Jangan mengajukan pra-registrasi sebelum seluruh temuan KRITIS terselesaikan."
            if is_id else
            f"DOSSIER SUBMISSION READINESS: {audit_result['status_text_en']}\n"
            f"AUDIT SUMMARY: Total {audit_result['total_gaps']} Findings | "
            f"Critical: {audit_result['critical_count']} | "
            f"Major: {audit_result['major_count']} | "
            f"Minor: {audit_result['minor_count']}\n"
            f"RECOMMENDATION: Do NOT file pre-registration until all CRITICAL findings are fully resolved."
        )
        sb_run.font.name = "Calibri"
        sb_run.font.size = Pt(10.5)
        sb_run.font.bold = True
        sb_run.font.color.rgb = border_text_color

        doc.add_paragraph().paragraph_format.space_after = Pt(12)

        # 2. Detailed Findings & Remediation Protocols Table
        doc.add_heading("Rincian Temuan Kesenjangan & Tindakan Remediasi" if is_id else "Detailed Gap Findings & Remediation Protocols", level=2)
        
        findings = audit_result["findings"]
        table_widths = [1.2, 1.8, 1.8, 2.2]
        tbl = doc.add_table(rows=len(findings) + 1, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        h = tbl.rows[0]
        h.cells[0].paragraphs[0].text = "Kode & Severity" if is_id else "Code & Severity"
        h.cells[1].paragraphs[0].text = "Modul & Temuan" if is_id else "Module & Observation"
        h.cells[2].paragraphs[0].text = "Dampak Regulasi BPOM" if is_id else "BPOM Impact"
        h.cells[3].paragraphs[0].text = "Protokol Tindakan Remediasi" if is_id else "Remediation Action Plan"
        style_row_as_header(h, table_widths)

        for i, item in enumerate(findings):
            row = tbl.rows[i + 1]
            sev = item["severity"]
            
            # Severity Cell
            p0 = row.cells[0].paragraphs[0]
            r0 = p0.add_run(f"[{sev}]\n{item['item_code']}")
            r0.font.bold = True
            r0.font.size = Pt(9)
            if sev == "CRITICAL":
                r0.font.color.rgb = COLOR_RED_CRITICAL
            elif sev == "MAJOR":
                r0.font.color.rgb = COLOR_AMBER_MAJOR
            else:
                r0.font.color.rgb = COLOR_MUTED

            # Module & Observation
            p1 = row.cells[1].paragraphs[0]
            r1_title = p1.add_run(f"{item['module']}\n{item['title_id'] if is_id else item['title_en']}\n\n")
            r1_title.font.bold = True
            r1_title.font.size = Pt(8.5)
            r1_obs = p1.add_run(item['observation_id'] if is_id else item['observation_en'])
            r1_obs.font.size = Pt(8.5)
            r1_obs.font.color.rgb = COLOR_DARK_TEXT

            # Impact & Basis
            p2 = row.cells[2].paragraphs[0]
            r2_imp = p2.add_run(f"{item['impact_id'] if is_id else item['impact_en']}\n\n")
            r2_imp.font.size = Pt(8.5)
            r2_basis = p2.add_run(f"Dasar Regulasi:\n{item['bpom_basis']}" if is_id else f"Regulatory Basis:\n{item['bpom_basis']}")
            r2_basis.font.bold = True
            r2_basis.font.size = Pt(8)
            r2_basis.font.color.rgb = COLOR_SECONDARY_BLUE

            # Remediation
            p3 = row.cells[3].paragraphs[0]
            r3_rem = p3.add_run(item['remediation_id'] if is_id else item['remediation_en'])
            r3_rem.font.size = Pt(8.5)
            r3_rem.font.color.rgb = COLOR_DARK_TEXT

            for c_idx, cell in enumerate(row.cells):
                cell.width = Inches(table_widths[c_idx])
                set_cell_margins(cell, 80, 80, 100, 100)
                if sev == "CRITICAL":
                    set_cell_background(cell, "FFF8F8")
                elif sev == "MAJOR":
                    set_cell_background(cell, "FFFCF5")
                elif i % 2 == 1:
                    set_cell_background(cell, HEX_LIGHT_GRAY)

        doc.add_paragraph().paragraph_format.space_after = Pt(16)

        # 3. Pre-Submission Sign-off Block
        doc.add_heading("Tanda Tangan Evaluator Regulasi" if is_id else "Regulatory Sign-Off Block", level=2)
        so_p = doc.add_paragraph()
        so_p.text = (
            "Diverifikasi dan Ditinjau Oleh:\n\n\n\n"
            "________________________________________\n"
            "Apoteker Penanggung Jawab Teknis (PJT)\n"
            "Departemen Regulatory Affairs & Quality Assurance"
            if is_id else
            "Verified and Reviewed By:\n\n\n\n"
            "________________________________________\n"
            "Lead Pharmacist / Head of Regulatory Affairs\n"
            "Regulatory Affairs & Quality Assurance Department"
        )
        so_p.runs[0].font.size = Pt(9.5)
        so_p.runs[0].font.color.rgb = COLOR_DARK_TEXT

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path

    @staticmethod
    def create_dossier_evaluation_report(evaluation_result: Any, product: Dict[str, Any],
                                         language: str, output_path: str):
        """Generates comprehensive BPOM Dossier Evaluation & Review Report."""
        doc = Document()
        is_id = (language.upper() == "ID")

        # Page setup
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

            header_p = section.header.paragraphs[0]
            header_p.text = "BADAN POM RI - DIREKTORAT REGISTRASI OBAT - LAPORAN EVALUASI DOSSIER"
            header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            header_p.runs[0].font.size = Pt(8.5)
            header_p.runs[0].font.color.rgb = COLOR_MUTED

            footer_p = section.footer.paragraphs[0]
            footer_p.text = "DOKUMEN RESMI EVALUASI TEKNIS DOSSIER REGISTRASI OBAT - RAHASIA & TERBATAS"
            footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            footer_p.runs[0].font.size = Pt(8)
            footer_p.runs[0].font.color.rgb = COLOR_MUTED

        # Main Title
        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(0)
        title_p.paragraph_format.space_after = Pt(4)
        run_title = title_p.add_run(
            "LAPORAN HASIL EVALUASI DOSSIER REGISTRASI OBAT" if is_id
            else "BPOM DRUG REGISTRATION DOSSIER EVALUATION & TECHNICAL ASSESSMENT"
        )
        run_title.font.name = "Calibri"
        run_title.font.size = Pt(19)
        run_title.font.bold = True
        run_title.font.color.rgb = COLOR_PRIMARY_NAVY

        # Subtitle
        sub_p = doc.add_paragraph()
        sub_p.paragraph_format.space_after = Pt(14)
        run_sub = sub_p.add_run(
            f"Evaluasi Berdasarkan Peraturan BPOM No. 24/2017 jo. PerBPOM 13/2021 & 24/2023 | Produk: {product.get('product_name', '-')}" if is_id
            else f"Technical Review under BPOM Regulation No. 24/2017 jo. 13/2021 & 24/2023 | Target: {product.get('product_name', '-')}"
        )
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(10)
        run_sub.font.color.rgb = COLOR_SECONDARY_BLUE

        # 1. Executive Verdict Banner
        v_code = evaluation_result.verdict_code
        if v_code == "MS":
            bg_hex = HEX_OK_BG
            text_color = COLOR_GREEN_OK
        elif v_code == "TD":
            bg_hex = HEX_MAJOR_BG
            text_color = COLOR_AMBER_MAJOR
        else:
            bg_hex = HEX_CRITICAL_BG
            text_color = COLOR_RED_CRITICAL

        callout = doc.add_table(rows=1, cols=1)
        callout.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_cell = callout.rows[0].cells[0]
        c_cell.width = Inches(7.0)
        set_cell_background(c_cell, bg_hex)
        set_cell_margins(c_cell, 120, 120, 150, 150)
        cp = c_cell.paragraphs[0]
        c_run = cp.add_run(
            f"KEPUTUSAN EVALUASI: {evaluation_result.verdict_title_id if is_id else evaluation_result.verdict_title_en}\n"
            f"SKOR KEPATUHAN DOSSIER: {evaluation_result.compliance_score}% / 100%\n"
            f"CATATAN TIM EVALUATOR: {evaluation_result.summary_remarks_id if is_id else evaluation_result.summary_remarks_en}"
        )
        c_run.font.name = "Calibri"
        c_run.font.size = Pt(10)
        c_run.font.bold = True
        c_run.font.color.rgb = text_color

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 2. Product Profile
        doc.add_heading("1. Data Identitas Obat & Pendaftaran" if is_id else "1. Medicinal Product & Submission Data", level=2)
        table = doc.add_table(rows=4, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        p_widths = [2.2, 4.8]
        rows_data = [
            ("Nama Produk Obat" if is_id else "Product Brand Name", product.get("product_name", "-")),
            ("Bentuk Sediaan & Kekuatan" if is_id else "Dosage Form & Strength", f"{product.get('dosage_form', '-')} | {product.get('strength', '-')}"),
            ("Klasifikasi & Asal" if is_id else "Classification & Origin", f"{product.get('product_type', '-')} ({product.get('origin', '-')})"),
            ("Tanggal Penilaian" if is_id else "Evaluation Date", "Tahun Anggaran 2026")
        ]
        for idx, (label, val) in enumerate(rows_data):
            row = table.rows[idx]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width, c1.width = Inches(p_widths[0]), Inches(p_widths[1])
            set_cell_background(c0, HEX_LIGHT_GRAY)
            set_cell_margins(c0, 60, 60, 100, 100)
            set_cell_margins(c1, 60, 60, 100, 100)
            c0.paragraphs[0].add_run(label).font.bold = True
            c0.paragraphs[0].runs[0].font.size = Pt(9)
            c1.paragraphs[0].add_run(str(val)).font.size = Pt(9)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 3. Module-by-Module Evaluation Table
        doc.add_heading("2. Hasil Penilaian Per Bagian Modul ACTD" if is_id else "2. Module-by-Module ACTD Assessment Findings", level=2)
        mod_results = evaluation_result.module_results
        t_widths = [1.0, 1.8, 0.9, 3.3]
        tbl = doc.add_table(rows=len(mod_results) + 1, cols=4)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Header
        h = tbl.rows[0]
        h.cells[0].paragraphs[0].text = "Kode" if is_id else "Code"
        h.cells[1].paragraphs[0].text = "Parameter Modul" if is_id else "Module Parameter"
        h.cells[2].paragraphs[0].text = "Status" if is_id else "Status"
        h.cells[3].paragraphs[0].text = "Catatan Hasil Penilaian Evaluator BPOM" if is_id else "Evaluator Observations & Findings"
        style_row_as_header(h, t_widths)

        for i, item in enumerate(mod_results):
            row = tbl.rows[i + 1]
            st = item["status"]
            
            row.cells[0].paragraphs[0].text = item["code"]
            row.cells[1].paragraphs[0].text = f"{item['module']}\n{item['title_id'] if is_id else item['title_en']}"
            row.cells[2].paragraphs[0].text = f"[{st}]\n{item['score']}/{item['max_score']}"
            row.cells[3].paragraphs[0].text = item["evaluator_notes_id"] if is_id else item["evaluator_notes_en"]

            for c_idx, cell in enumerate(row.cells):
                cell.width = Inches(t_widths[c_idx])
                set_cell_margins(cell, 60, 60, 80, 80)
                if st == "TMS":
                    set_cell_background(cell, "FFF5F5")
                elif st == "TD":
                    set_cell_background(cell, "FFFAF0")
                elif i % 2 == 1:
                    set_cell_background(cell, HEX_LIGHT_GRAY)
                cell.paragraphs[0].runs[0].font.size = Pt(8.5)
                if c_idx == 2:
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell.paragraphs[0].runs[0].font.bold = True
                    if st == "MS":
                        cell.paragraphs[0].runs[0].font.color.rgb = COLOR_GREEN_OK
                    elif st == "TD":
                        cell.paragraphs[0].runs[0].font.color.rgb = COLOR_AMBER_MAJOR
                    elif st == "TMS":
                        cell.paragraphs[0].runs[0].font.color.rgb = COLOR_RED_CRITICAL

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 4. Draft Formal Surat Permintaan Tambahan Data BPOM
        doc.add_heading("3. Konsep Surat Permintaan Tambahan Data BPOM RI" if is_id else "3. Official BPOM Request for Additional Data (Query Letter)", level=2)
        td_box = doc.add_table(rows=1, cols=1)
        td_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        td_c = td_box.rows[0].cells[0]
        td_c.width = Inches(7.0)
        set_cell_background(td_c, "F8FAFC")
        set_cell_margins(td_c, 140, 140, 160, 160)
        
        td_p = td_c.paragraphs[0]
        td_p.paragraph_format.line_spacing = 1.15
        letter_content = evaluation_result.draft_letter_id if is_id else evaluation_result.draft_letter_en
        td_run = td_p.add_run(letter_content)
        td_run.font.name = "Consolas"
        td_run.font.size = Pt(8.5)
        td_run.font.color.rgb = COLOR_DARK_TEXT

        doc.add_paragraph().paragraph_format.space_after = Pt(16)

        # 5. Evaluator Sign-off
        doc.add_heading("4. Tanda Tangan Tim Penilai Khasiat, Keamanan & Mutu" if is_id else "4. Regulatory Evaluators Sign-Off", level=2)
        so_p = doc.add_paragraph()
        so_p.text = (
            "Diverifikasi oleh Penilai Mutu & Klinis BPOM RI:\n\n\n\n"
            "________________________________________\n"
            "Tim Penilai Registrasi Obat / Kasubdit Evaluasi\n"
            "Direktorat Registrasi Obat - Badan Pengawas Obat dan Makanan RI"
            if is_id else
            "Verified by Technical & Clinical Evaluators:\n\n\n\n"
            "________________________________________\n"
            "Lead Regulatory Reviewer / Head of Drug Assessment\n"
            "Directorate of Drug Registration - Badan POM RI"
        )
        so_p.runs[0].font.size = Pt(9.5)
        so_p.runs[0].font.color.rgb = COLOR_DARK_TEXT

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        doc.save(output_path)
        return output_path

